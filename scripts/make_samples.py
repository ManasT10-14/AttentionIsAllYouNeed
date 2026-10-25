"""Generate demo contracts (DOCX, authoring spec §3.1) with deliberately planted traps.
Run:  python scripts/make_samples.py      → corpus/samples/contracts/*.docx
Replace these with the teammate's real documents when ready (same format)."""
from pathlib import Path

from docx import Document

OUT = Path(__file__).resolve().parents[1] / "corpus" / "samples" / "contracts"

CONTRACTS = {
    "01_office_lease.docx": ({
        "CONTRACT TITLE": "Office Lease Agreement", "CONTRACT ID": "LEASE-2025-001", "CONTRACT TYPE": "lease",
        "EFFECTIVE DATE": "2025-01-01", "PARTY A": 'Acme Analytics Pvt Ltd (the "Tenant")',
        "PARTY B": 'Skyline Realty LLP (the "Landlord")', "GOVERNING LAW": "India", "VERSION": "1"}, [
        "1. DEFINITIONS",
        '1.1 "Premises" means Unit 7B, Prestige Tech Park, Bengaluru, measuring 8,000 sq. ft.',
        "2. TERM AND RENEWAL",
        "2.1 The lease commences on the Effective Date and continues for a period of 24 months.",
        "2.2 The lease shall renew automatically for successive periods of 12 months unless the Tenant gives the "
        "Landlord written notice of non-renewal at least 60 days before the expiry of the then-current term.",
        "2.3 On each renewal the monthly rent shall increase by 10% over the rent payable immediately before renewal.",
        "3. RENT AND PAYMENT",
        "3.1 The Tenant shall pay a monthly rent of INR 4,50,000 in advance on or before the 5th day of each month, "
        "the first payment being due on 5 January 2025.",
        "3.2 Any rent not paid when due shall carry interest at 18% per annum from the due date until payment.",
        "3.3 The Tenant has paid a refundable security deposit of INR 27,00,000.",
        "4. MAINTENANCE",
        "4.1 The Landlord shall carry out structural repairs within a reasonable time after notice from the Tenant.",
        "4.2 The Tenant shall bear all costs of internal repairs, fit-outs and common-area maintenance charges as "
        "determined by the Landlord from time to time.",
        "5. TERMINATION",
        "5.1 The Landlord may terminate this lease by giving the Tenant 30 days written notice.",
        "5.2 The Tenant may terminate this lease only by giving the Landlord 6 months written notice and paying "
        "rent for the notice period in advance.",
        "5.3 If the Tenant terminates within the first 12 months of any term, the security deposit is forfeited.",
        "6. LIABILITY AND INDEMNITY",
        "6.1 The Tenant shall indemnify the Landlord against all losses, claims and expenses arising from the use of "
        "the Premises, without limit.",
        "7. GOVERNING LAW AND JURISDICTION",
        "7.1 This lease is governed by the laws of India and the courts at Mumbai shall have exclusive jurisdiction.",
    ]),
    "02_nimbus_cloud_msa.docx": ({
        "CONTRACT TITLE": "Master Cloud Services Agreement", "CONTRACT ID": "MSA-2026-001", "CONTRACT TYPE": "saas",
        "EFFECTIVE DATE": "2026-02-01", "PARTY A": 'Acme Analytics Pvt Ltd (the "Customer")',
        "PARTY B": 'Nimbus Cloud Ltd (the "Provider")', "GOVERNING LAW": "Singapore", "VERSION": "1"}, [
        "1. SERVICES",
        "1.1 The Provider shall provide managed cloud hosting for the Customer's analytics platform, which the "
        "Customer uses to deliver services to its own customers.",
        "2. TERM AND RENEWAL",
        "2.1 This Agreement starts on the Effective Date and continues for 12 months.",
        "2.2 This Agreement renews automatically for further 12-month periods unless either party gives written "
        "notice of non-renewal at least 30 days before the end of the current term.",
        "3. FEES",
        "3.1 The Customer shall pay monthly fees of INR 3,00,000, invoiced on the 1st of each month starting "
        "1 February 2026.",
        "3.2 The Provider may revise the fees at any time on 30 days notice by email.",
        "4. PAYMENT",
        "4.1 Late payments attract interest at 2% per month on the overdue amount.",
        "4.2 Invoices are payable within 30 days of the invoice date.",
        "4.3 If any amount remains unpaid for more than 45 days the Provider may suspend the Services without "
        "further notice.",
        "5. SERVICE LEVELS",
        "5.1 The Provider shall ensure monthly availability of 99.5%.",
        "5.2 For each 0.5% (or part) by which availability falls below 99.5%, the Customer is entitled to a "
        "service credit of 5% of the monthly fee, capped at 10% of the monthly fee.",
        "5.3 Service credits are the Customer's sole and exclusive remedy for any failure to meet service levels.",
        "6. LIMITATION OF LIABILITY",
        "6.1 The Provider's total liability under this Agreement shall not exceed the fees paid by the Customer in "
        "the 3 months preceding the claim.",
        "7. DATA PROTECTION",
        "7.1 The Provider may process Customer data in any country where it or its sub-processors operate.",
        "8. GOVERNING LAW",
        "8.1 This Agreement is governed by the laws of Singapore and disputes shall be referred to arbitration "
        "seated in Singapore.",
        "SCHEDULE A - PAYMENT TERMS",
        "All invoices are payable within 15 days of the invoice date.",
    ]),
    "03_bigretail_customer_agreement.docx": ({
        "CONTRACT TITLE": "Customer Services Agreement", "CONTRACT ID": "CSA-2026-004", "CONTRACT TYPE": "saas",
        "EFFECTIVE DATE": "2026-04-01", "PARTY A": 'Acme Analytics Pvt Ltd (the "Provider")',
        "PARTY B": 'BigRetail India Ltd (the "Client")', "GOVERNING LAW": "India", "VERSION": "1"}, [
        "1. SERVICES",
        "1.1 The Provider shall deliver its retail analytics platform to the Client, operated by the Provider's "
        "team from its premises at Prestige Tech Park, Bengaluru.",
        "2. TERM",
        "2.1 This Agreement runs for 24 months from the Effective Date.",
        "3. FEES",
        "3.1 The Client shall pay the Provider INR 12,00,000 per month, payable on the 10th day of each month "
        "starting 10 April 2026.",
        "4. SERVICE LEVELS",
        "4.1 The Provider guarantees monthly platform availability of 99.9%.",
        "4.2 For each month in which availability is below 99.9%, the Provider shall pay the Client a penalty of 10% "
        "of the monthly fee, without any cap.",
        "4.3 Three consecutive months below the guaranteed availability entitle the Client to terminate "
        "immediately and claim damages.",
        "5. TERMINATION",
        "5.1 The Client may terminate for convenience on 30 days notice.",
        "6. GOVERNING LAW",
        "6.1 This Agreement is governed by the laws of India; courts at Bengaluru have exclusive jurisdiction.",
    ]),
    "04_orion_supply_agreement.docx": ({
        "CONTRACT TITLE": "Exclusive Supply Agreement", "CONTRACT ID": "SUP-2025-007", "CONTRACT TYPE": "supply",
        "EFFECTIVE DATE": "2025-04-01", "PARTY A": 'Acme Analytics Pvt Ltd (the "Buyer")',
        "PARTY B": 'Orion Components Pvt Ltd (the "Supplier")', "GOVERNING LAW": "India", "VERSION": "1"}, [
        "1. SUPPLY",
        "1.1 The Supplier shall supply IoT footfall sensors to the Buyer.",
        "2. EXCLUSIVITY",
        "2.1 During the term the Buyer shall purchase IoT sensors exclusively from the Supplier and shall not "
        "procure equivalent products from any third party.",
        "3. TERM",
        "3.1 This Agreement continues for 36 months from the Effective Date.",
        "4. MINIMUM PURCHASE",
        "4.1 The Buyer shall purchase sensors worth at least INR 20,00,000 in each calendar quarter.",
        "4.2 If purchases in a quarter fall below the minimum, the Buyer shall pay the Supplier 15% of the shortfall "
        "within 30 days of the quarter end.",
        "5. TERMINATION",
        "5.1 Either party may terminate for material breach not cured within 30 days of written notice.",
        "6. GOVERNING LAW",
        "6.1 This Agreement is governed by the laws of India.",
    ]),
    "05_vega_supply_agreement.docx": ({
        "CONTRACT TITLE": "Supply Agreement", "CONTRACT ID": "SUP-2026-002", "CONTRACT TYPE": "supply",
        "EFFECTIVE DATE": "2026-06-01", "PARTY A": 'Acme Analytics Pvt Ltd (the "Buyer")',
        "PARTY B": 'Vega Electronics Ltd (the "Seller")', "GOVERNING LAW": "India", "VERSION": "1"}, [
        "1. PRODUCTS",
        "1.1 The Seller shall supply IoT footfall sensors to the Buyer against purchase orders.",
        "2. DELIVERY",
        "2.1 The Seller shall deliver each order within 14 days of the purchase order.",
        "2.2 For late delivery the Seller shall pay liquidated damages at the rate set out in Schedule B.",
        "3. PRICE AND PAYMENT",
        "3.1 The Buyer shall pay INR 6,00,000 per quarter for the committed volume, payable on the first day of "
        "each quarter starting 1 July 2026.",
        "4. TERM AND TERMINATION",
        "4.1 This Agreement continues for 12 months from the Effective Date.",
        "4.2 The Buyer may terminate on 60 days notice.",
        "4.3 The Buyer may terminate on 30 days notice.",
        "5. GOVERNING LAW",
        "5.1 This Agreement is governed by the laws of India.",
    ]),
    "06_bigretail_nda.docx": ({
        "CONTRACT TITLE": "Mutual Non-Disclosure Agreement", "CONTRACT ID": "NDA-2026-003", "CONTRACT TYPE": "nda",
        "EFFECTIVE DATE": "2026-03-15", "PARTY A": 'Acme Analytics Pvt Ltd (the "Company")',
        "PARTY B": 'BigRetail India Ltd (the "Counterparty")', "GOVERNING LAW": "England and Wales", "VERSION": "1"}, [
        "1. CONFIDENTIALITY",
        "1.1 Each party shall keep the other's Confidential Information secret for 5 years from disclosure.",
        "1.2 On request, the receiving party shall return or destroy Confidential Information within a reasonable time.",
        "2. REMEDIES",
        "2.1 A breach of this Agreement entitles the disclosing party to liquidated damages of INR 50,00,000 per breach.",
        "3. GOVERNING LAW",
        "3.1 This Agreement is governed by the laws of England and Wales and the English courts have exclusive "
        "jurisdiction.",
    ]),
}


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for name, (cover, clauses) in CONTRACTS.items():
        d = Document()
        for k, v in cover.items():
            d.add_paragraph(f"{k}: {v}")
        d.add_paragraph("")
        for c in clauses:
            d.add_paragraph(c)
        d.save(OUT / name)
        print("wrote", OUT / name)


if __name__ == "__main__":
    main()
