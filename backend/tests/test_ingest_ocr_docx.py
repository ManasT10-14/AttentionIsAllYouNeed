"""OCR (scanned PDF / image) and DOCX ingestion produce clauses, pages and highlight boxes."""
import fitz
from docx import Document

from odin.ingest.parse import docx_to_pdf, image_to_pdf, parse_pdf

CLAUSES = ["1. PAYMENT", "1.1 The Customer shall pay all invoices within 30 days.",
           "2. TERMINATION", "2.1 Either party may terminate on 60 days notice."]


def _scanned_pdf(tmp_path):
    """Render text to a bitmap and embed it as an image-only page (no text layer)."""
    src = fitz.open()
    pg = src.new_page()
    for i, line in enumerate(CLAUSES):
        pg.insert_text((50, 80 + 40 * i), line, fontsize=20)
    png = pg.get_pixmap(dpi=150).tobytes("png")
    out = fitz.open()
    page = out.new_page(width=595, height=842)
    page.insert_image(page.rect, stream=png)
    path = tmp_path / "scan.pdf"
    out.save(path)
    return path, png


def test_scanned_pdf_is_ocrd(tmp_path):
    path, _ = _scanned_pdf(tmp_path)
    assert not fitz.open(path)[0].get_text().strip()
    doc = parse_pdf(path, "SCAN")
    assert any("recovered by OCR" in w and w.startswith("Page 1") for w in doc.warnings)
    assert "30 days" in doc.text and "60 days" in doc.text
    c = next(c for c in doc.clauses if c.number.startswith("1"))
    assert c.page == 1 and c.bboxes and "invoices" in doc.text


def test_image_upload_becomes_pdf(tmp_path):
    _, png = _scanned_pdf(tmp_path)
    dst = tmp_path / "img.pdf"
    image_to_pdf(png, ".png", dst)
    assert "PAYMENT" in parse_pdf(dst, "IMG").text.upper()


def test_docx_without_libreoffice(tmp_path, monkeypatch):
    monkeypatch.setattr("odin.ingest.parse._soffice", lambda: None)
    d = Document()
    d.add_paragraph("CONTRACT ID: DOCX-1")
    d.add_paragraph("PAYMENT", style="List Number")
    d.add_paragraph("Invoices are due in 45 days.")
    d.add_table(rows=1, cols=2).rows[0].cells[0].text = "Fee"
    src = tmp_path / "a.docx"
    d.save(src)
    pdf = tmp_path / "a.pdf"
    assert docx_to_pdf(src, pdf) == []      # python-docx succeeded: no scary warning
    doc = parse_pdf(pdf, "A")
    assert "45 days" in doc.text and "Fee" in doc.text
