"""Apply an approved amendment: redlines → new DOCX version → re-ingest → update graph (closes the loop)."""
from __future__ import annotations

import io

from odin.core.context import current
from odin.db import repo


def build_docx(contract_id: str, redlines: dict[str, str], *, source_version: int | None = None,
               source_sha256: str | None = None) -> bytes:
    from docx import Document
    from odin.services.approvals import ApprovalRequired

    c = repo.contract(contract_id)
    if source_version is not None:
        version = repo.version_row(contract_id, source_version)
        if (not c or c.current_version != source_version or c.status in {"withdrawn", "deleted", "superseded"}
            or not version or version.sha256 != source_sha256):
            raise ApprovalRequired("amendment source changed before reading its approved version")
    meta = dict(c.meta)
    doc = Document()
    cover = [("CONTRACT TITLE", meta.get("title") or c.title), ("CONTRACT ID", contract_id),
             ("CONTRACT TYPE", meta.get("type", "other")), ("EFFECTIVE DATE", meta.get("effective_date") or "")]
    for i, p in enumerate(meta.get("parties", [])):
        alias = f' (the "{p["alias"]}")' if p.get("alias") else ""
        cover.append((f"PARTY {chr(65 + i)}", p["name"] + alias))
    cover += [("GOVERNING LAW", meta.get("governing_law", "")), ("VERSION", str(c.current_version + 1)),
              ("SUPERSEDES", f"{contract_id} v{c.current_version}")]
    for k, v in cover:
        doc.add_paragraph(f"{k}: {v}")
    doc.add_paragraph("")
    for cl in repo.clauses(contract_id, version=source_version, current_only=source_version is None):
        text = redlines.get(cl.number, cl.text)
        if not text.lstrip().startswith(cl.number.split("#")[0]):
            text = f"{cl.number} {text}"
        for line in text.split("\n"):
            doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def apply(action_id: str, *, expected_sources: list[dict] | None = None) -> dict:
    from odin.ingest.service import ingest_contract
    from odin.services.approvals import ApprovalRequired, dispatch_sources

    a = repo.get("action", action_id)
    if not a or a.get("kind") != "amendment" or not a.get("redlines"):
        raise ValueError("not an amendment with redlines")
    sources = dispatch_sources(action_id, "apply_amendment", {"action_id": action_id})
    if expected_sources != sources:
        raise ApprovalRequired("amendment requires its approved source snapshot")
    source = next((item for item in sources if item["contract_id"] == a["contract_id"]), None)
    if not source:
        raise ApprovalRequired("approved amendment source is unavailable")
    red = {r["clause_number"]: r["after"] for r in a["redlines"]}
    company_id = current().company_id
    content = build_docx(a["contract_id"], red, source_version=source["version"], source_sha256=source["sha256"])
    # Do not ingest a proposal built from a version that changed during document construction.
    dispatch_sources(action_id, "apply_amendment", {"action_id": action_id})
    res = ingest_contract(company_id, f"{a['contract_id']}_amended.docx", content,
                          approval_execution_id=a["approval_execution"]["id"])
    repo.set_status(action_id, "applied", applied_version=res.version)
    # The API launches analysis after registry.finish releases this action's durable lease.
    return {"contract_id": res.contract_id, "version": res.version, "previous_version": res.previous_version,
            "is_update": res.is_update, "run_id": None, "warnings": res.warnings}
