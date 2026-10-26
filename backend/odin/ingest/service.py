"""Upload → files on disk → parsed version → clause rows (+embeddings). Detects new vs updated contracts."""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path

from sqlalchemy import select, update

from odin.core.config import DATA_DIR, odin_config
from odin.core.hooks import HOOKS
from odin.core.models import CompanyContext, new_id
from odin.db.session import db
from odin.db.tables import ClauseRow, Company, Contract, ContractVersion, Record
from odin.ingest import pack
from odin.ingest.parse import docx_to_pdf, image_to_pdf, parse_money, parse_pdf, text_to_pdf
from odin.retrieval.hybrid import chunk_header, embed_passages


@dataclass
class IngestResult:
    contract_id: str
    version: int
    is_update: bool
    previous_version: int | None
    clause_count: int
    warnings: list[str]
    duplicate: bool = False


IMAGE_EXTS = (".png", ".jpg", ".jpeg", ".tif", ".tiff")


def _safe(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", name)[:80]


def ingest_contract(company_id: str, filename: str, content: bytes, *, twin_pdf: bytes | None = None,
                    doc_id: str | None = None, title: str | None = None) -> IngestResult:
    """doc_id forces the Document ID (bundle: filename stem); otherwise cover CONTRACT ID, else the stem.
    A dev-pack .txt is the canonical citation text; twin_pdf (or a rendering of the txt) feeds the viewer."""
    ext = Path(filename).suffix.lower()
    if ext not in (".pdf", ".docx", ".txt", ".md", *IMAGE_EXTS):
        raise ValueError(f"Unsupported file type {ext} — use PDF, DOCX or a scanned image (PNG/JPG/TIFF).")
    sha = hashlib.sha256(content).hexdigest()
    stage = DATA_DIR / company_id / "_staging" / sha[:16]
    stage.mkdir(parents=True, exist_ok=True)
    src = stage / f"original{ext}"
    src.write_bytes(content)
    pdf = stage / "document.pdf"
    warnings: list[str] = []
    canonical = None
    if ext in (".txt", ".md"):
        canonical = content.decode("utf-8-sig", errors="replace").replace("\r\n", "\n").replace("\r", "\n")
        if not pack.is_pack(canonical):
            canonical = None
    if canonical is not None:
        if twin_pdf:
            pdf.write_bytes(twin_pdf)
        else:
            text_to_pdf(canonical, pdf)
            warnings.append("No PDF twin — viewer shows a rendering of the canonical text.")
    elif ext == ".pdf":
        pdf.write_bytes(content)
    elif ext == ".docx":
        warnings += docx_to_pdf(src, pdf)
    elif ext in IMAGE_EXTS:
        image_to_pdf(content, ext, pdf)
    else:
        text_to_pdf(content.decode("utf-8", errors="replace"), pdf)

    fallback = _safe(Path(filename).stem).upper()
    if canonical is not None:
        doc, doc.extra = pack.parse_pack(canonical, pdf, doc_id, fallback, title or "")
    else:
        doc = parse_pdf(pdf, fallback_id=fallback)
        if doc_id:
            doc.cover.contract_id = doc_id
        if title:
            doc.cover.title = title
    HOOKS.on_ingest(doc.text, {"company_id": company_id, "filename": filename})
    warnings += doc.warnings
    cover = doc.cover
    cid = cover.contract_id.strip()

    with db() as s:
        existing = s.scalar(select(Contract).where(Contract.company_id == company_id, Contract.contract_id == cid))
        if existing:
            last = s.scalar(select(ContractVersion).where(ContractVersion.company_id == company_id,
                                                          ContractVersion.contract_id == cid,
                                                          ContractVersion.version == existing.current_version))
            if last and last.sha256 == sha:
                return IngestResult(cid, existing.current_version, False, None, 0, ["Identical file already ingested."],
                                    duplicate=True)
            prev = existing.current_version
            version = max(prev + 1, cover.version or 0)
        else:
            prev = None
            deleted = s.scalars(select(Record).where(Record.company_id == company_id,
                                 Record.contract_id == cid, Record.kind == "contract_deletion")).all()
            last_deleted = max((r.data.get("last_version", 0) for r in deleted), default=0)
            version = max(last_deleted + 1, cover.version or 1)

    final = DATA_DIR / company_id / _safe(cid) / f"v{version}"

    headers = [chunk_header(cid, version, c.number, c.heading) + "\n" + c.text[:1800] for c in doc.clauses]
    vectors = (embed_passages(headers) if odin_config()["retrieval"].get("dense_enabled", True)
               else [None] * len(headers)) if headers else []

    cover_json = {**cover.model_dump(mode="json"), **getattr(doc, "extra", {})}
    cover_json["version"] = version
    expected_id = existing.id if existing else None
    with db() as s:
        s.scalar(select(Company).where(Company.id == company_id).with_for_update())
        current = s.scalar(select(Contract).where(Contract.company_id == company_id, Contract.contract_id == cid))
        if (current.id if current else None) != expected_id or (current and current.current_version != prev):
            raise ValueError("Contract changed or was deleted during upload; retry the upload")
        final.mkdir(parents=True, exist_ok=True)
        (final / src.name).write_bytes(src.read_bytes())
        (final / "document.pdf").write_bytes(pdf.read_bytes())
        if prev is None:
            s.add(Contract(id=new_id("ctr"), company_id=company_id, contract_id=cid, title=cover.title or cid,
                           type=cover.type, current_version=version, meta=cover_json, status="ingested"))
        else:
            s.execute(update(Contract).where(Contract.company_id == company_id, Contract.contract_id == cid)
                      .values(current_version=version, meta=cover_json, title=cover.title or cid,
                              type=cover.type, status="updated"))
            s.execute(update(ClauseRow).where(ClauseRow.company_id == company_id, ClauseRow.contract_id == cid)
                      .values(is_current=False))
        s.add(ContractVersion(id=new_id("ver"), company_id=company_id, contract_id=cid, version=version,
                              file_path=str(final / src.name), pdf_path=str(final / "document.pdf"), sha256=sha,
                              text=doc.text, cover=cover_json, parse_warnings=warnings))
        for c, v in zip(doc.clauses, vectors):
            s.add(ClauseRow(id=f"{company_id}:{cid}:v{version}:{c.number}", company_id=company_id, contract_id=cid,
                            version=version, is_current=True, number=c.number, heading=c.heading, text=c.text,
                            ref=getattr(c, "ref", None), parent_ref=getattr(c, "parent_ref", None),
                            clause_type=c.clause_type, page=c.page, char_start=c.char_start,
                            char_end=c.char_end, bboxes=c.bboxes, embedding=v))
    return IngestResult(cid, version, prev is not None, prev, len(doc.clauses), warnings)


# ───────────────────────────── company context ─────────────────────────────
_CTX_KEYS = {"COMPANY NAME": "name", "INDUSTRY": "industry", "HQ / JURISDICTIONS": "jurisdictions",
             "JURISDICTIONS": "jurisdictions", "FISCAL YEAR START": "fiscal_year_start", "BASE CURRENCY": "base_currency",
             "RISK APPETITE": "risk_appetite", "APPROVAL THRESHOLD": "approval_threshold", "KEY CONTACTS": "contacts",
             "STRATEGIC PRIORITIES": "priorities", "CRITICAL DEPENDENCIES": "critical_deps"}


def _sections(text: str) -> dict[str, list[str]]:
    """ALL-CAPS heading underlined with dashes → its lines (dev-pack company/context.txt)."""
    out: dict[str, list[str]] = {"": []}
    cur = ""
    lines = text.splitlines()
    for i, ln in enumerate(lines):
        if i + 1 < len(lines) and re.fullmatch(r"[A-Z][A-Z0-9 &/'-]{2,60}(?:\([a-z ]+\))?", ln.strip()) and re.fullmatch(r"-{3,}", lines[i + 1].strip()):
            cur = ln.strip().upper()
            out[cur] = []
        elif not re.fullmatch(r"[-=]{3,}", ln.strip()):
            out[cur].append(ln)
    return out


def parse_pack_context(text: str) -> CompanyContext:
    """Dev-pack context.txt: LEGAL NAME / JURISDICTION header, signatories, notice addresses, insurance."""
    from odin.core.models import InsurancePolicy
    from odin.ingest.parse import parse_date
    sec = _sections(text)
    head = {m.group(1).strip().upper(): m.group(2).strip()
            for m in re.finditer(r"^([A-Z][A-Z ]+):\s+(.+)$", "\n".join(sec[""]), re.M)}
    contacts: list[str] = []
    sig = sec.get("AUTHORIZED SIGNATORIES", [])
    for i, ln in enumerate(sig):
        m = re.match(r"^\s*\d+\.\s+(.+?),\s+(.+)$", ln)
        if m:
            mail = next((re.search(r"\S+@\S+", x).group(0) for x in sig[i + 1:i + 3] if "@" in x), "")
            contacts.append(f"{m.group(1)}, {m.group(2)} (authorized signatory) <{mail}>")
    for ln in sec.get("KEY PERSONNEL (OPERATIONS)", []):
        m = re.match(r"^\s*(.+?),\s+(.+?)\s+[—–-]\s+(\S+@\S+)", ln)
        if m:
            contacts.append(f"{m.group(1)}, {m.group(2)} <{m.group(3)}>")
    notice, blocks, cur = "", [], None
    for ln in sec.get("NOTICE ADDRESSES", []):
        if re.match(r"^For .+:$", ln.strip()):
            cur = [ln.strip().rstrip(":")]
            blocks.append(cur)
        elif cur is not None and ln.strip():
            cur.append(ln.strip().replace("Email: ", ""))
    for b in blocks:
        if "contractual" in b[0].lower() or "legal" in b[0].lower():
            notice = ", ".join(b[1:])
        else:
            contacts.append(f"{b[0]}: " + ", ".join(b[1:]))
    insurance, pol = [], None
    keys = {"policy": "type", "insurer": "insurer", "policy no.": "policy_number", "coverage": "coverage",
            "effective": "effective", "expiry": "expiry"}
    for ln in sec.get("INSURANCE", []):
        m = re.match(r"^\s*([A-Za-z. ]+?):\s+(.+)$", ln)
        if not m or m.group(1).lower() not in keys:
            continue
        k, v = keys[m.group(1).lower()], m.group(2).strip()
        if k == "type":
            pol = {}
            insurance.append(pol)
        if pol is not None:
            pol[k] = parse_date(v) if k in ("effective", "expiry") else v
    desc = " ".join(" ".join(sec.get("BUSINESS DESCRIPTION", [])).split())
    sentences = re.split(r"(?<=\.)\s+(?=[A-Z])", desc)
    if head.get("PRINCIPAL OFFICE"):
        contacts.insert(0, f"Principal office: {head['PRINCIPAL OFFICE']}")
    fy = parse_date((head.get("FISCAL YEAR") or "").split("–")[0].split("-")[0] + " 2000")
    juris = head.get("JURISDICTION", "")
    return CompanyContext(name=head.get("LEGAL NAME", ""), industry=sentences[0][:300],
                          critical_deps=" ".join(x for x in sentences if "relies on" in x or "third-party" in x),
                          jurisdictions=[juris] if juris else [],
                          base_currency="USD" if "United States" in juris else "INR",
                          fiscal_year_start=fy.strftime("%m-%d") if fy else "01-01", contacts=contacts,
                          insurance=[InsurancePolicy(**p) for p in insurance], notice_address=notice, raw_text=text)


def parse_company_context(text: str) -> CompanyContext:
    if re.search(r"^LEGAL NAME:", text, re.M):
        return parse_pack_context(text)
    vals: dict = {"raw_text": text}
    for line in text.splitlines():
        m = re.match(r"^\s*([A-Z][A-Z /]+?)\s*:\s*(.+?)\s*$", line)
        if not m or m.group(1).strip().upper() not in _CTX_KEYS:
            continue
        key, val = _CTX_KEYS[m.group(1).strip().upper()], m.group(2)
        if key in ("jurisdictions", "contacts"):
            vals[key] = [v.strip() for v in re.split(r"[;,]", val) if v.strip()]
        elif key == "approval_threshold":
            vals[key] = parse_money(val)
        elif key == "risk_appetite":
            v = val.strip().lower()
            vals[key] = v if v in ("low", "medium", "high") else "medium"
        else:
            vals[key] = val.strip()
    return CompanyContext(**vals)


def ingest_company_context(company_id: str, filename: str, content: bytes) -> CompanyContext:
    if filename.lower().endswith(".pdf"):
        import fitz
        with fitz.open(stream=content, filetype="pdf") as d:
            text = "\n".join(p.get_text() for p in d)
    elif filename.lower().endswith(".docx"):
        import io
        from docx import Document
        text = "\n".join(p.text for p in Document(io.BytesIO(content)).paragraphs)
    else:
        text = content.decode("utf-8", errors="replace")
    text = HOOKS.on_ingest(text, {"company_id": company_id, "filename": filename, "kind": "company_context"})
    ctx = parse_company_context(text)
    with db() as s:
        c = s.get(Company, company_id)
        c.context = ctx.model_dump(mode="json")
        if ctx.name:
            c.name = ctx.name
    return ctx
